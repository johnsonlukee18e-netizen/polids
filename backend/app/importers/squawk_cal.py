"""SQUAWK: pule kodów SSR z arkusza EUROCONTROL / ICAO EUR CAL (SSR Code Allocation List, załącznik do ICAO EUR
Doc 023) → data/seed/squawk.json.

    python -m backend.app.importers.squawk_cal data/import/eurocontrol-icao-eur-cal-v4-13.xls   # zapis do data/seed/squawk.json
    python -m backend.app.importers.squawk_cal data/import/eurocontrol-icao-eur-cal-v4-13.xls --dry-run   # tylko raport

Arkusz czytamy tylko jako dane (bez makr i formuł): .xls przez xlrd, .xlsx przez openpyxl. To narzędzia
z requirements-dev.txt, w obrazie aplikacji ich nie ma (importer uruchamia się raz, na komputerze dewelopera).

Układ CAL (rozpoznawany wprost, bez zgadywania): arkusz z nagłówkiem
    Series | From | To | PA | Unit | CIV/MIL | APP | Unit which assigned | A/R | Destination | Via | Protection | Remarks | Ver
Series, From i To to po dwie cyfry ósemkowe (w arkuszu liczba albo tekst: 75, 1, 77 albo "75", "01", "77"),
kod = Series + From/To: 75 / 01 / 77 → 7501–7577, 13 / 00 / 27 → 1300–1327. PA = obszar uczestniczący (EUR-A…EUR-J,
CCAMS, NATO), Unit = prefiksy ICAO / organy, które używają kodów w tym obszarze, A/R: A = kody przydziela Unit,
R = retencja (kody przydzielone przez "Unit which assigned" zostają po wlocie do obszaru). Arkusz "PAs": obszar →
państwa / organy. Wiersz z cyfrą 8/9 albo odwróconym zakresem jest wypisywany i pomijany.

Wybór (Polska = Unit EP, obszar EUR-E), każda pula z kategorią:
    epww   PA EUR-E, A, Unit = dokładnie EP – pule, z których FIR EPWW sam przydziela kody
    ccams  PA CCAMS, A, Unit zawiera EP – kody przydzielane centralnie przez CCAMS (także dla Polski)
    nato   PA NATO, A, Unit zawiera EP – pule wojskowe NATO
    eur-e  pozostałe wiersze PA EUR-E – przydziały innych członków EUR-E (A) i retencja (R)

Inny układ arkusza: zapasowo ogólne rozpoznawanie kolumn (wiersz nagłówka z UNIT, role kolumn zgadywane z nazw
i treści, filtr UNIT EP / EP… i wzmianka EUR-E; --generic wymusza ten tryb, --*-col wskazuje kolumny ręcznie).

Wynik: pools [{from, to, count, category, pa, unit, civ_mil, app, assigned_by, ar, destination, via, protection,
remarks, ver, sheet, row}] posortowane po kategorii i kodzie, meta {source, layout, edition, generated,
participating_areas, counts…}. Kody specjalne ("special") i opis pliku ("_uwaga") zostają z dotychczasowego
pliku; notatka "_status" (stara: import oczekuje) znika."""

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from ..config import settings
from ..services.squawk import CATEGORIES, code_value, parse_codes

COMMAND = "python -m backend.app.importers.squawk_cal"
HOME_UNIT = "EP"      # Polska w kolumnie Unit
HOME_PA = "EUR-E"     # obszar uczestniczący, do którego należy Polska
UNITS = ["EP"]        # filtr układu ogólnego
MENTIONS = ["EUR-E"]

# układ CAL: nazwa kolumny w arkuszu → klucz puli; bez kolumn wymaganych arkusz nie jest w układzie CAL
CAL_COLUMNS = [("Series", "series"), ("From", "from"), ("To", "to"), ("PA", "pa"), ("Unit", "unit"),
               ("CIV/MIL", "civ_mil"), ("APP", "app"), ("Unit which assigned", "assigned_by"), ("A/R", "ar"),
               ("Destination", "destination"), ("Via", "via"), ("Protection", "protection"), ("Remarks", "remarks"),
               ("Ver", "ver")]
CAL_REQUIRED = {"series", "from", "to", "pa", "unit", "ar"}
CAL_LISTS = ("unit", "assigned_by", "destination", "via", "protection")   # listy prefiksów ICAO / obszarów
PA_NAME = re.compile(r"^[A-Z]{2,6}(?:-[A-Z0-9]{1,3})?$")
VER = re.compile(r"^V?\s*(\d{1,2}\.\d{1,3})$", re.I)

# układ ogólny: nazwy kolumn (po angielsku jak w arkuszu); kolumny kodów rozpoznajemy też po treści
CODE_NAME = re.compile(r"CODE|SSR|SQUAWK|SERIES|BLOCK|RANGE", re.I)
FROM_NAME = re.compile(r"\b(FROM|START|FIRST|BEGIN|LOW|LOWER)\b", re.I)
TO_NAME = re.compile(r"\b(TO|END|LAST|UPPER|HIGH|UNTIL)\b", re.I)
TYPE_NAME = re.compile(r"TYPE|USE|USAGE|PURPOSE|CATEGORY|ALLOC|ASSIGN|STATUS|KIND", re.I)
REMARKS_NAME = re.compile(r"REMARK|NOTE|COMMENT", re.I)
SKIP_NAME = re.compile(r"^(NO|NR|#|ROW|ITEM|SERIAL|S/N|LP)\.?$|DATE|EDITION|AMDT|AMENDMENT|VALID", re.I)
EDITION = re.compile(r"\b(?:EDITION|ED\.|VERSION|VER\.?|V)\s*(\d{1,2}\.\d{1,3})\b", re.I)
EDITION_FILE = re.compile(r"v(\d{1,2})[-_.](\d{1,3})(?!\d)", re.I)


class CalError(Exception):
    """Błąd, który importer pokazuje zamiast śladu stosu (brak biblioteki, zły plik, brak kolumny)."""


def _text(v) -> str:
    """Komórka → napis bez spacji na brzegach; liczby całkowite bez ".0" (xlrd daje 100.0 dla kodu 0100)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return re.sub(r"[ \t\r\xa0]+", " ", str(v)).strip()


def _fill(grid: list[list[str]], r0: int, r1: int, c0: int, c1: int) -> None:
    """Scalone komórki [r0, r1) × [c0, c1): puste dostają wartość lewej górnej."""
    if r0 >= len(grid):
        return
    top = grid[r0][c0] if c0 < len(grid[r0]) else ""
    for r in range(r0, min(r1, len(grid))):
        for c in range(c0, min(c1, len(grid[r]))):
            if not grid[r][c]:
                grid[r][c] = top


def read_book(path: Path) -> dict[str, list[list[str]]]:
    """{arkusz: wiersze napisów} z wypełnionymi scalonymi komórkami; biblioteki importowane dopiero tu."""
    suffix = path.suffix.lower()
    sheets = {}
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise CalError("Brak modułu xlrd (pliki .xls): pip install -r requirements-dev.txt "
                           "albo pip install xlrd") from exc
        try:
            try:
                book = xlrd.open_workbook(str(path), formatting_info=True)
            except NotImplementedError:   # część plików .xls nie ma informacji o formatowaniu (bez scaleń)
                book = xlrd.open_workbook(str(path))
        except Exception as exc:  # noqa: BLE001 - xlrd rzuca różne wyjątki dla uszkodzonych plików
            raise CalError(f"{path.name}: nie da się odczytać arkusza .xls ({exc})") from exc
        for sh in book.sheets():
            grid = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    v = sh.cell_value(r, c)
                    if sh.cell_type(r, c) == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate_as_datetime(v, book.datemode).date()
                        except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                            pass
                    row.append(_text(v))
                grid.append(row)
            for r0, r1, c0, c1 in getattr(sh, "merged_cells", None) or []:
                _fill(grid, r0, r1, c0, c1)
            for row in grid:   # formatting_info: puste, ale sformatowane kolumny (w CAL do 256)
                while row and not row[-1]:
                    row.pop()
            sheets[sh.name] = grid
    elif suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as exc:
            raise CalError("Brak modułu openpyxl (pliki .xlsx): pip install -r requirements-dev.txt "
                           "albo pip install openpyxl") from exc
        try:
            book = openpyxl.load_workbook(path, read_only=False, data_only=True)   # read_only nie zna scaleń
        except Exception as exc:  # noqa: BLE001 - openpyxl rzuca różne wyjątki dla uszkodzonych plików
            raise CalError(f"{path.name}: nie da się odczytać arkusza .xlsx ({exc})") from exc
        for ws in book.worksheets:
            grid = [[_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
            for rng in ws.merged_cells.ranges:
                _fill(grid, rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col)
            sheets[ws.title] = grid
    else:
        raise CalError(f"{path.name}: obsługiwane są pliki .xls i .xlsx")
    return sheets


# --- układ EUROCONTROL CAL

def _head(cell: str) -> str:
    return re.sub(r"\s+", " ", cell).strip().upper()


def cal_layout(grid: list[list[str]], scan: int = 10) -> tuple[int, dict[str, int]] | None:
    """(indeks wiersza nagłówka, {klucz: kolumna}), gdy arkusz ma nagłówek CAL (nazwy kolumn dokładnie jak
    w CAL_COLUMNS, wielkość liter bez znaczenia) z wszystkimi kolumnami wymaganymi; inaczej None."""
    names = {_head(n): k for n, k in CAL_COLUMNS}
    for i, row in enumerate(grid[:scan]):
        cols = {}
        for j, cell in enumerate(row):
            if (k := names.get(_head(cell))) and k not in cols:
                cols[k] = j
        if CAL_REQUIRED <= set(cols):
            return i, cols
    return None


def octal2(text: str) -> str | None:
    """Series / From / To: 1–2 cyfry ósemkowe → dwie ("1" z liczby 1.0 → "01"); None: 8/9, tekst, pusto."""
    return text.zfill(2) if re.fullmatch(r"[0-7]{1,2}", text) else None


def tokens(text: str) -> list[str]:
    """Lista prefiksów ICAO / organów z komórki ("EP EY  lA" → ["EP", "EY", "LA"])."""
    return text.upper().split()


def ver_text(text: str) -> str:
    """Wersja CAL z kolumny Ver: "V4.12" i "4.12" → "4.12"."""
    m = VER.match(text.strip())
    return m.group(1) if m else text.strip()


def ver_key(ver: str) -> tuple[int, int]:
    m = re.fullmatch(r"(\d+)\.(\d+)", ver)
    return (int(m.group(1)), int(m.group(2))) if m else (-1, -1)


def category(pa: str, ar: str, units: list[str]) -> str | None:
    """Kategoria wiersza CAL (patrz opis modułu) albo None (wiersz nas nie dotyczy)."""
    if pa == HOME_PA and ar == "A" and units == [HOME_UNIT]:
        return "epww"
    if pa == "CCAMS" and ar == "A" and HOME_UNIT in units:
        return "ccams"
    if pa == "NATO" and ar == "A" and HOME_UNIT in units:
        return "nato"
    if pa == HOME_PA:
        return "eur-e"
    return None


def read_pas(grid: list[list[str]]) -> dict[str, list[str]]:
    """Arkusz "PAs": obszar uczestniczący (kolumna 1) → państwa / organy (kolumna 2, lista po spacjach)."""
    out = {}
    for row in grid:
        name = _head(row[0]) if row else ""
        if PA_NAME.match(name) and len(row) > 1 and tokens(row[1]):
            out[name] = tokens(row[1])
    return out


def extract_cal(sheets: dict[str, list[list[str]]], layouts: dict[str, tuple[int, dict[str, int]]], source: str,
                report=print) -> dict:
    """Pule z arkuszy w układzie CAL (layouts: wynik cal_layout dla tych arkuszy): {meta, pools, skipped}."""
    pools, skipped, vers, areas, ar_bad, rows_total = [], [], [], {}, [], 0
    for name, grid in sheets.items():
        if name not in layouts and name.strip().upper() in ("PAS", "PA"):
            areas.update(read_pas(grid))
            report(f"\n## arkusz {name!r}: obszary uczestniczące (PA)")
            for pa, members in read_pas(grid).items():
                report(f"   {pa:<6} {' '.join(members)}")
    for name, (hi, cols) in layouts.items():
        grid = sheets[name]
        report(f"\n## arkusz {name!r}: układ EUROCONTROL CAL, nagłówek w wierszu {hi + 1}: "
               f"{' | '.join(grid[hi][j] for j in sorted(cols.values()))}")
        missing = [n for n, k in CAL_COLUMNS if k not in cols]
        if missing:
            report(f"   brak kolumn (puste w pulach): {', '.join(missing)}")
        pa_ar, rows = Counter(), 0
        for i, raw in enumerate(grid[hi + 1:], start=hi + 2):
            cell = {k: (raw[j] if j < len(raw) else "") for k, j in cols.items()}
            if not any(cell.values()):
                continue
            rows += 1
            if v := cell.get("ver"):
                vers.append(ver_text(v))
            pa, ar = _head(cell["pa"]), _head(cell["ar"])
            pa_ar[(pa, ar)] += 1
            s, f, t = (octal2(cell[k]) for k in ("series", "from", "to"))
            why = None
            if s is None or f is None or t is None:
                why = "Series/From/To: dwie cyfry ósemkowe 0–7"
            elif code_value(s + f) > code_value(s + t):
                why = "From > To"
            if why:
                skipped.append({"sheet": name, "row": i, "why": why, "pa": pa,
                                "raw": {n: cell[k] for n, k in CAL_COLUMNS if cell.get(k)}})
                continue
            if ar not in ("A", "R"):
                ar_bad.append(f"{name}:{i} {pa} A/R={cell['ar']!r}")
            units = tokens(cell["unit"])
            cat = category(pa, ar, units)
            if cat is None:
                continue
            pool = {"from": s + f, "to": s + t, "count": code_value(s + t) - code_value(s + f) + 1, "category": cat,
                    "pa": pa, "unit": " ".join(units), "civ_mil": _head(cell.get("civ_mil", "")),
                    "app": _head(cell.get("app", "")), "assigned_by": " ".join(tokens(cell.get("assigned_by", ""))),
                    "ar": ar}
            for k in ("destination", "via", "protection"):
                pool[k] = " ".join(tokens(cell.get(k, "")))
            pool["remarks"] = re.sub(r"\s*\n\s*", " / ", cell.get("remarks", "")).strip()
            pool["ver"] = ver_text(cell.get("ver", ""))
            pool["sheet"], pool["row"] = name, i
            pools.append(pool)
        rows_total += rows
        report(f"   wiersze danych: {rows}")
        report("   PA × A/R: " + ", ".join(f"{pa or '–'} {ar or '–'} {n}" for (pa, ar), n in sorted(pa_ar.items())))
    if ar_bad:
        report(f"   A/R inne niż A i R ({len(ar_bad)}): {'; '.join(ar_bad)}")
    order = {c: k for k, c in enumerate(CATEGORIES)}
    pools.sort(key=lambda p: (order[p["category"]], code_value(p["from"]), code_value(p["to"]), p["sheet"], p["row"]))
    counts = {c: sum(p["category"] == c for p in pools) for c in CATEGORIES}
    latest = max(vers, key=ver_key, default=None)
    latest = latest if latest and ver_key(latest) >= (0, 0) else None
    named = f"{m.group(1)}.{m.group(2)}" if (m := EDITION_FILE.search(source)) else None
    edition = named or latest
    edition_from = " and ".join(w for w, v in (("file name", named), ("Ver column", latest)) if v == edition and v)
    if named and latest and named != latest:
        report(f"   UWAGA: wydanie z nazwy pliku {named}, a najnowsza zmiana wiersza (Ver) {latest}")
    if areas and HOME_UNIT not in areas.get(HOME_PA, []):
        report(f"   UWAGA: w arkuszu PAs obszar {HOME_PA} nie zawiera {HOME_UNIT}")
    meta = {"source": source, "layout": "CAL", "edition": edition, "edition_from": edition_from or None,
            "latest_ver": latest, "generated": date.today().isoformat(), "unit": HOME_UNIT, "pa": HOME_PA,
            "participating_areas": areas, "counts": counts, "rows": rows_total, "skipped": len(skipped)}
    return {"meta": meta, "pools": pools, "skipped": skipped}


# --- układ ogólny (zapasowo, gdy arkusz nie ma nagłówka CAL)

def header_row(grid: list[list[str]], scan: int = 60) -> int | None:
    """Indeks wiersza nagłówka: pierwsza komórka równa UNIT, a gdy takiej nie ma – zawierająca słowo UNIT."""
    for i, row in enumerate(grid[:scan]):
        if any(c.upper() == "UNIT" for c in row):
            return i
    for i, row in enumerate(grid[:scan]):
        if any(re.search(r"\bUNIT\b", c, re.I) for c in row):
            return i
    return None


def columns(grid: list[list[str]], hi: int, width: int) -> tuple[list[str], int]:
    """Nazwy kolumn i indeks pierwszego wiersza danych. Nagłówek dwupoziomowy (scalone "CODES" nad "FROM"/"TO"):
    powtórzone nazwy łączymy z podpisami z następnego wiersza, o ile to tekst, a nie kody."""
    head = [re.sub(r"\s+", " ", h).strip() for h in (grid[hi] + [""] * width)[:width]]
    start = hi + 1
    dup = {n for n, k in Counter(head).items() if n and k > 1}
    if dup and start < len(grid):
        sub = [re.sub(r"\s+", " ", c).strip() for c in (grid[start] + [""] * width)[:width]]
        if any(sub[j] and head[j] in dup for j in range(width)) and \
                all(not sub[j] or parse_codes(sub[j]) is None for j in range(width)):
            head = [f"{h} {s}".strip() if s and s != h else h for h, s in zip(head, sub)]
            start += 1
    names, seen = [], Counter()
    for j, h in enumerate(head):
        name = h or f"col{j + 1}"
        seen[name] += 1
        names.append(name if seen[name] == 1 else f"{name}#{seen[name]}")
    return names, start


def _find(names: list[str], wanted: str | None, role: str) -> int | None:
    if wanted is None:
        return None
    for j, n in enumerate(names):
        if n.upper() == wanted.strip().upper():
            return j
    raise CalError(f"--{role}-col {wanted!r}: nie ma takiej kolumny (są: {', '.join(names)})")


def detect_roles(names: list[str], rows: list[list[str]], forced: dict | None = None) -> dict:
    """Role kolumn: unit, from + to (dwie kolumny pojedynczych kodów) albo code (zakresy w jednej kolumnie),
    type (przeznaczenie), remarks (uwagi). forced: nazwy kolumn z wiersza poleceń (pierwszeństwo)."""
    forced = forced or {}
    roles = {role: _find(names, forced.get(role), role) for role in ("unit", "from", "to", "code", "type", "remarks")}
    if roles["unit"] is None:
        roles["unit"] = next((j for j, n in enumerate(names) if n.upper() == "UNIT"),
                             next((j for j, n in enumerate(names) if re.search(r"\bUNIT\b", n, re.I)), None))

    def share(j: int) -> tuple[float, float]:
        vals = [r[j] for r in rows if j < len(r) and r[j]]
        if not vals:
            return 0.0, 0.0
        ok = sum(parse_codes(v) is not None for v in vals)
        # pojedynczy kod: 3–4 cyfry (0100 zapisane liczbą to 100); 1–2 cyfry to raczej numer serii albo wiersza
        single = sum(bool(re.fullmatch(r"[0-7]{3,4}", v)) for v in vals)
        return ok / len(vals), single / len(vals)

    if roles["from"] is None and roles["code"] is None:
        taken = {roles["unit"], roles["type"], roles["remarks"]}
        coded, singles = [], []
        for j, n in enumerate(names):
            if j in taken or SKIP_NAME.search(n):
                continue
            hinted = bool(CODE_NAME.search(n) or FROM_NAME.search(n) or TO_NAME.search(n))
            ok, single = share(j)
            if ok >= 0.9 or (hinted and ok >= 0.5):
                coded.append((hinted, ok, j))
                if single >= 0.9 or (hinted and single >= 0.5):
                    singles.append(j)
        if len(singles) >= 2:
            first = next((j for j in singles if FROM_NAME.search(names[j])), singles[0])
            rest = [j for j in singles if j != first]
            roles["from"] = first
            roles["to"] = next((j for j in rest if TO_NAME.search(names[j])), next((j for j in rest if j > first), rest[0]))
        elif coded:
            roles["code"] = max(coded, key=lambda x: (x[0], x[1], -x[2]))[2]
    if roles["from"] is not None and roles["to"] is None and roles["code"] is None:
        nxt = roles["from"] + 1
        if nxt < len(names) and TO_NAME.search(names[nxt]):
            roles["to"] = nxt
    taken = {v for v in roles.values() if v is not None}
    if roles["type"] is None:
        roles["type"] = next((j for j, n in enumerate(names) if j not in taken and TYPE_NAME.search(n)), None)
        taken.add(roles["type"])
    if roles["remarks"] is None:
        roles["remarks"] = next((j for j, n in enumerate(names) if j not in taken and REMARKS_NAME.search(n)), None)
    return roles


def unit_match(unit: str, units: list[str]) -> bool:
    """UNIT równy celowi albo (cel dwuliterowy, np. EP) zaczynający się od niego; kilka organów w komórce
    (EPWW/EPKK, przecinki) sprawdzamy osobno."""
    tokens = [t for t in re.split(r"[\s/,;+&]+", unit.upper()) if t]
    return any(t == u or (len(u) == 2 and t.startswith(u)) for t in tokens for u in (x.upper() for x in units))


def mention_pattern(mentions: list[str]) -> re.Pattern | None:
    """EUR-E jako osobne słowo (nie EUR-EAST, nie XEUR-E), wielkość liter bez znaczenia."""
    if not mentions:
        return None
    return re.compile("|".join(rf"(?<![A-Z0-9]){re.escape(m)}(?![A-Z0-9])" for m in mentions), re.I)


def find_edition(sheets: dict[str, list[list[str]]], heads: dict[str, int | None], source: str) -> tuple[str | None, str | None]:
    """Wydanie CAL: z nazw arkuszy i wierszy nad nagłówkiem (albo pierwszych 15 wierszy), inaczej z nazwy pliku
    (eurocontrol-icao-eur-cal-v4-13.xls → 4.13)."""
    for name, grid in sheets.items():
        top = grid[:heads[name]] if heads.get(name) is not None else grid[:15]
        for text in [name] + [c for row in top for c in row if c]:
            if m := EDITION.search(text):
                return m.group(1), "sheet"
    if m := EDITION_FILE.search(source):
        return f"{m.group(1)}.{m.group(2)}", "file name"
    return None, None


def extract_generic(sheets: dict[str, list[list[str]]], source: str, units: list[str] | None = None,
                    mentions: list[str] | None = None, forced: dict | None = None, report=print,
                    max_values: int = 40) -> dict:
    """Pule z arkuszy w nieznanym układzie (role kolumn zgadywane): {meta, pools, skipped}. Kategoria: epww, gdy
    pasuje UNIT, inaczej eur-e (sama wzmianka EUR-E)."""
    units = UNITS if units is None else units
    mentions = MENTIONS if mentions is None else mentions
    patterns = {m: mention_pattern([m]) for m in mentions}
    pools, skipped, info, heads = [], [], {}, {}
    for name, grid in sheets.items():
        width = max((len(r) for r in grid), default=0)
        hi = header_row(grid)
        heads[name] = hi
        report(f"\n## arkusz {name!r}: wiersze {len(grid)}, kolumny {width}")
        if hi is None:
            report("   brak wiersza nagłówka z kolumną UNIT – pomijam")
            continue
        names, start = columns(grid, hi, width)
        rows = [((r + [""] * width)[:width], start + k + 1) for k, r in enumerate(grid[start:])]
        rows = [(r, n) for r, n in rows if any(r)]
        roles = detect_roles(names, [r for r, _ in rows], forced)
        role_names = {k: names[v] for k, v in roles.items() if v is not None}
        report(f"   nagłówek: wiersz {hi + 1}" + (f" (+ podpisy w wierszu {hi + 2})" if start > hi + 1 else ""))
        above = [[c for c in r if c] for r in grid[:hi] if any(r)]
        if above:
            report(f"   nad nagłówkiem: {above}")
        report(f"   kolumny: {names}")
        report(f"   role: {role_names}")
        report(f"   wiersze danych: {len(rows)}")
        for j, col in enumerate(names):
            cnt = Counter(r[j] for r, _ in rows if r[j])
            shown = dict(cnt.most_common(max_values)) if len(cnt) <= max_values else \
                f"{len(cnt)} różnych, np. {list(cnt)[:12]}"
            report(f"   [{j + 1}] {col}: {shown}")
        eur = Counter(c for r, _ in rows for c in r if c and "EUR" in c.upper())
        if eur:
            report(f"   komórki z EUR: {dict(eur.most_common(max_values))}")
        info[name] = {"header_row": hi + 1, "columns": role_names}
        if roles["unit"] is None:
            report("   brak kolumny UNIT – pomijam")
            continue
        code_cols = [roles[k] for k in ("from", "to", "code") if roles[k] is not None]
        for r, rn in rows:
            why = []
            if unit_match(r[roles["unit"]], units):
                why.append("unit")
            why.extend(m for m, pat in patterns.items() if any(pat.search(c) for c in r if c))
            if not why:
                continue
            raw = {names[j]: (r[j].zfill(4) if j in code_cols and re.fullmatch(r"[0-7]{1,3}", r[j]) else r[j])
                   for j in range(width) if r[j]}
            if roles["from"] is not None:
                first = parse_codes(r[roles["from"]])
                last = parse_codes(r[roles["to"]]) if roles["to"] is not None and r[roles["to"]] else first
                ranges = [(first[0][0], last[-1][1])] if first and last and len(first) == len(last) == 1 and \
                    code_value(first[0][0]) <= code_value(last[-1][1]) else None
            elif roles["code"] is not None:
                ranges = parse_codes(r[roles["code"]])
            else:
                ranges = None
            if not ranges:
                skipped.append({"sheet": name, "row": rn, "match": why, "raw": raw})
                continue
            for a, b in ranges:
                pools.append({"from": a, "to": b, "count": code_value(b) - code_value(a) + 1,
                              "category": "epww" if "unit" in why else "eur-e",
                              "unit": r[roles["unit"]],
                              "type": r[roles["type"]] if roles["type"] is not None else "",
                              "remarks": r[roles["remarks"]] if roles["remarks"] is not None else "",
                              "match": why, "sheet": name, "row": rn, "raw": raw})
    pools.sort(key=lambda p: (code_value(p["from"]), code_value(p["to"]), p["unit"], p["sheet"], p["row"]))
    edition, edition_from = find_edition(sheets, heads, source)
    meta = {"source": source, "layout": "generic", "edition": edition, "edition_from": edition_from,
            "generated": date.today().isoformat(), "filter": {"unit": units, "mention": mentions}, "sheets": info,
            "counts": {c: sum(p["category"] == c for p in pools) for c in CATEGORIES},
            "pools": len(pools), "skipped": len(skipped)}
    return {"meta": meta, "pools": pools, "skipped": skipped}


def extract(sheets: dict[str, list[list[str]]], source: str, units: list[str] | None = None,
            mentions: list[str] | None = None, only: list[str] | None = None, forced: dict | None = None,
            report=print, max_values: int = 40, generic: bool = False) -> dict:
    """Pule z arkuszy (dict jak z read_book): {meta, pools, skipped}. Arkusze z nagłówkiem CAL → układ CAL
    (units / mentions / forced go nie dotyczą); bez nich albo z generic=True → układ ogólny."""
    sheets = {k: v for k, v in sheets.items() if not only or k in only or (k.strip().upper() in ("PAS", "PA"))}
    layouts = {} if generic else {k: lay for k, v in sheets.items() if (lay := cal_layout(v))}
    if layouts:
        if units is not None or mentions is not None or forced:
            report("UWAGA: --unit, --mention i --*-col dotyczą tylko układu ogólnego (--generic); pomijam")
        return extract_cal(sheets, layouts, source, report)
    sheets = {k: v for k, v in sheets.items() if not only or k in only}
    return extract_generic(sheets, source, units, mentions, forced, report, max_values)


def _tree(v, ind: int) -> str:
    """Słownik wcięty (klucz na wiersz), wszystko inne (też listy) w jednym wierszu."""
    if isinstance(v, dict) and v:
        pad = " " * ind
        return "{\n" + ",\n".join(f"{pad} {json.dumps(k, ensure_ascii=False)}: {_tree(x, ind + 1)}" for k, x in v.items()) \
            + f"\n{pad}}}"
    return json.dumps(v, ensure_ascii=False)


def dumps(data: dict) -> str:
    """Tekst squawk.json: klucze główne po jednym na wiersz, meta wcięta (listy w jednym wierszu), jedna pula /
    jeden kod na wiersz (czytelne zmiany w gicie)."""
    def value(v) -> str:
        if isinstance(v, list):
            return "[" + ",".join(f"\n  {json.dumps(x, ensure_ascii=False)}" for x in v) + ("\n ]" if v else "]")
        return _tree(v, 1)
    return "{\n" + ",\n".join(f" {json.dumps(k)}: {value(v)}" for k, v in data.items()) + "\n}\n"


def merge(base: dict, result: dict) -> dict:
    """Nowa treść pliku: opis i kody specjalne z dotychczasowego pliku, meta i pule z arkusza, bez "_status"."""
    out = {k: v for k, v in base.items() if k not in ("_status", "meta", "special", "pools")}
    out["meta"] = result["meta"]
    out["special"] = base.get("special") or []
    out["pools"] = result["pools"]
    return out


def _base(out: Path) -> dict:
    for path in (out, settings.seed_dir / "squawk.json"):
        try:
            data = json.loads(path.read_text("utf-8"))
            if isinstance(data, dict):
                return data
        except (OSError, ValueError):
            continue
    return {}


def _print_pool(p: dict) -> None:
    if "pa" in p:   # układ CAL
        print(f"   {p['from']}–{p['to']} ({p['count']:>3})  {p['pa']:<5} {p['ar'] or '–'} {p['civ_mil'] or '–':<7} "
              f"{p['unit'][:28]:<28} przydziela {p['assigned_by'][:20] or '–':<20} cel {p['destination'][:16] or '–':<16} "
              f"ochrona {p['protection'] or '–'} | {p['remarks'] or '–'} | V{p['ver']}   [{p['sheet']}:{p['row']}]")
    else:
        print(f"   {p['from']}–{p['to']} ({p['count']:>4})  {p['unit']:<10} {'+'.join(p['match']):<10} "
              f"{p['type'] or '–'} | {p['remarks'] or '–'}   [{p['sheet']}:{p['row']}]")


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(prog=COMMAND, description="Pule kodów SSR z arkusza EUROCONTROL CAL → squawk.json")
    ap.add_argument("file", type=Path, help="arkusz CAL (.xls albo .xlsx)")
    ap.add_argument("--out", type=Path, default=settings.seed_dir / "squawk.json", help="plik wynikowy (domyślnie data/seed/squawk.json)")
    ap.add_argument("--sheet", action="append", help="tylko te arkusze (powtarzalne)")
    ap.add_argument("--generic", action="store_true", help="układ ogólny (zgadywanie kolumn) zamiast układu CAL")
    ap.add_argument("--unit", action="append", help="układ ogólny: UNIT do wybrania (powtarzalne; domyślnie EP = EP i EP…)")
    ap.add_argument("--mention", action="append", help="układ ogólny: tekst w dowolnej kolumnie (powtarzalne; domyślnie EUR-E)")
    for role, what in (("unit", "organ"), ("from", "pierwszy kod puli"), ("to", "ostatni kod puli"),
                       ("code", "kody/zakresy w jednej kolumnie"), ("type", "przeznaczenie"), ("remarks", "uwagi")):
        ap.add_argument(f"--{role}-col", dest=f"col_{role}", metavar="NAZWA", help=f"układ ogólny, kolumna: {what}")
    ap.add_argument("--max-values", type=int, default=40, help="układ ogólny: ile różnych wartości kolumny wypisać")
    ap.add_argument("--dry-run", action="store_true", help="tylko raport, bez zapisu pliku")
    args = ap.parse_args(argv)
    forced = {role: getattr(args, f"col_{role}") for role in ("unit", "from", "to", "code", "type", "remarks")
              if getattr(args, f"col_{role}")}
    try:
        if not args.file.is_file():
            raise CalError(f"{args.file}: nie ma takiego pliku")
        sheets = read_book(args.file)
        print(f"# {args.file.name}: arkusze {list(sheets)}")
        result = extract(sheets, args.file.name, args.unit, args.mention, args.sheet, forced,
                         max_values=args.max_values, generic=args.generic)
    except CalError as exc:
        print(f"BŁĄD: {exc}", file=sys.stderr)
        return 2
    meta = result["meta"]
    print(f"\n# układ: {meta['layout']}")
    print(f"# wydanie CAL: {meta['edition'] or 'nie znaleziono'}" + (f" ({meta['edition_from']})" if meta["edition"] else ""))
    if meta["layout"] == "CAL":
        print(f"# wybór: Unit {meta['unit']}, obszar {meta['pa']} ({' '.join(meta['participating_areas'].get(meta['pa'], [])) or '?'})")
    else:
        print(f"# filtr: UNIT {meta['filter']['unit']} (dwuliterowy = prefiks), wzmianka {meta['filter']['mention']}")
    print(f"# pule: {len(result['pools'])} (" + ", ".join(f"{c} {n}" for c, n in meta["counts"].items()) + ")")
    for c in CATEGORIES:
        group = [p for p in result["pools"] if p["category"] == c]
        if group:
            print(f"\n## {c}: {len(group)}")
            for p in group:
                _print_pool(p)
    if result["skipped"]:
        print(f"\n# wiersze pominięte (złe kody): {len(result['skipped'])}")
        for s in result["skipped"]:
            print(f"   [{s['sheet']}:{s['row']}] {s.get('why') or '+'.join(s.get('match', []))} {s['raw']}")
    if args.dry_run:
        print("\n--dry-run: bez zapisu")
        return 0
    data = merge(_base(args.out), result)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(dumps(data), encoding="utf-8", newline="\n")
    print(f"\nzapisano {args.out}: {len(result['pools'])} pul, {len(data['special'])} kodów specjalnych")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
